# -*- coding: utf-8 -*-
"""Fifth pass: certificate page re-laid in the department format (members left, faculty guide right, countersigned head centred)."""
import copy, os
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(os.path.dirname(os.path.dirname(HERE)), "UW-TIG_Project_Work_Report_final.docx")
d = docx.Document(DOC)
ps = d.paragraphs
c0 = [i for i, p in enumerate(ps) if p.text.strip() == "CERTIFICATE"][0]
body_p = ps[c0 + 2]
assert body_p.text.startswith("Certified that this report titled")
end_p = [p for p in ps[c0:] if p._p.xpath("./w:pPr/w:sectPr") and p.text.strip() == ""][0]   # section-break paragraph after the certificate

# remove everything between the body paragraph and the section-break paragraph
node = body_p._p.getnext()
while node is not end_p._p:
    nxt = node.getnext()
    node.getparent().remove(node)
    node = nxt

NAMES = [("Akhil Ramalingam", "23Z207"), ("Anbuchandiran K", "23Z209"), ("Neelesh Padmanabh", "23Z241"),
         ("Saumiyaa Sri V L", "23Z261"), ("Sowndarya Elangi S", "23Z269"), ("Therdhana J P", "23Z273")]


def fmt(run, bold=False, size=12):
    run.font.name = "Times New Roman"
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


# ---- certificate paragraph: bold title and names, as in the department format
for r in list(body_p.runs):
    r._r.getparent().remove(r._r)
pieces = [("Certified that this report titled “", 0), ("UW-TIG: Uncertainty-Weighted Temporal Information Gain for Closed-Loop Defect Inspection", 1),
          ("” for the Project Work I (23Z711) is a bonafide work of ", 0)]
for k, (n, rn) in enumerate(NAMES):
    pieces.append((f"{n} ({rn})", 1))
    pieces.append((", " if k < len(NAMES) - 1 else " ", 0))
pieces.append(("who have carried out the work under my supervision for the partial fulfillment of the requirements for the award of the degree of "
               "Bachelor of Engineering in Computer Science and Engineering. Certified further that to the best of my knowledge and belief, the work "
               "reported herein does not form part of any other thesis or dissertation based on which a degree or an award was conferred on an earlier occasion.", 0))
for t, b in pieces:
    fmt(body_p.add_run(t), bold=bool(b))
body_p.alignment = AL.JUSTIFY
body_p.paragraph_format.line_spacing = 1.4


def new_par(text="", bold=False, align=AL.LEFT, before=None, space_after=0, size=12):
    el = OxmlElement("w:p")
    end_p._p.addprevious(el)
    p = Paragraph(el, end_p._parent)
    p.style = d.styles["Normal"]
    if text:
        fmt(p.add_run(text), bold=bold, size=size)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(space_after), 1.15
    pf.left_indent = pf.first_line_indent = None
    return p


new_par("")
new_par("Place: Coimbatore")
new_par("")
new_par("Date: 02.11.2026")
new_par("")

# ---- two-column block: members (left) and faculty guide / co-guide (right)
tbl = d.add_table(rows=1, cols=2)
end_p._p.addprevious(tbl._tbl)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
for c, w in zip(tbl.rows[0].cells, (Cm(7.2), Cm(8.8))):
    c.width = w


def cell_lines(cell, lines):
    first = True
    for text, bold, align in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.style = d.styles["Normal"]
        fmt(p.add_run(text), bold=bold)
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.left_indent = p.paragraph_format.first_line_indent = None


cell_lines(tbl.rows[0].cells[0], [("Members", True, AL.LEFT)] + [(n, False, AL.LEFT) for n, _ in NAMES])
cell_lines(tbl.rows[0].cells[1], [
    ("Faculty Guide", True, AL.RIGHT), ("Ms. L. Karthika", False, AL.RIGHT),
    ("Department of Computer Science and Engineering", False, AL.RIGHT), ("PSG College of Technology", False, AL.RIGHT), ("Coimbatore - 641004", False, AL.RIGHT),
    ("", False, AL.RIGHT),
    ("Co-Guide", True, AL.RIGHT), ("Dr. N. Gopika Rani", False, AL.RIGHT),
    ("Department of Computer Science and Engineering", False, AL.RIGHT), ("PSG College of Technology", False, AL.RIGHT), ("Coimbatore - 641004", False, AL.RIGHT)])
# borderless
tblPr = tbl._tbl.tblPr
borders = OxmlElement("w:tblBorders")
for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
    e = OxmlElement(f"w:{b}")
    e.set(qn("w:val"), "nil")
    borders.append(e)
tblPr.append(borders)

new_par("")
new_par("")
new_par("COUNTERSIGNED", align=AL.CENTER, space_after=14)
new_par("Dr. G. R. Karpagam", align=AL.CENTER)
new_par("HEAD", align=AL.CENTER)
new_par("Department of Computer Science and Engineering", align=AL.CENTER)
new_par("PSG College of Technology", align=AL.CENTER)
new_par("Coimbatore – 641004", align=AL.CENTER)
d.save(DOC)
print("saved")
