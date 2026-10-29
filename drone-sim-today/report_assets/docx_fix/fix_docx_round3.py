# -*- coding: utf-8 -*-
"""Third pass: Fig. 4.5 replaced by a clean flow-chart, the garbled two-step-lookahead paragraph rebuilt with
real subscripts, the blank page after the Introduction objectives removed, and the appendices dropped."""
import io
import os

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
DOC = os.path.join(ROOT, "UW-TIG_Project_Work_Report_final.docx")
FIG = os.path.join(HERE, "fig45_decision_flow.png")
d = docx.Document(DOC)
ps = d.paragraphs

# ---- 1. Fig. 4.5 picture
cap_idx = [i for i, p in enumerate(ps) if p.text.strip() == "Fig. 4.5 UW-TIG decision flow for one step"][0]
for j in range(cap_idx - 4, cap_idx):
    blips = ps[j]._p.xpath(".//a:blip")
    if blips:
        part = d.part.related_parts[blips[0].get(qn("r:embed"))]
        ext = os.path.splitext(part.partname)[1].lower()
        buf = io.BytesIO()
        Image.open(FIG).convert("RGB").save(buf, "JPEG" if ext in (".jpg", ".jpeg") else "PNG", quality=95)
        part._blob = buf.getvalue()
        print("Fig 4.5 image replaced in", part.partname)
        break

# ---- 2. Section 4.9.3 paragraph (garbled by the PDF-to-Word conversion)
start = [i for i, p in enumerate(ps) if p.text.startswith("In the full-utility phase, each candidate")][0]
end = [i for i, p in enumerate(ps) if i > start and p.text.startswith("current uncertainty, growth and staleness")][0]
main = ps[start]
for q in ps[start + 1:end + 1]:
    q._p.getparent().remove(q._p)
for r in list(main.runs):
    r._r.getparent().remove(r._r)
pieces = [("In the full-utility phase, each candidate v", 0), ("1", 1), (" is scored as U(v", 0), ("1", 1), (") + 0.5·max", 0), ("v2", 1),
          (" U(v", 0), ("2", 1), (" | after v", 0), ("1", 1),
          ("), where the second step uses the view counts that would result from visiting v", 0), ("1", 1),
          (" but the current uncertainty, growth and staleness (a future detection cannot be simulated). Only the first step is executed, "
           "after which the planner re-plans, in the manner of a receding-horizon planner [24].", 0)]
for text, sub in pieces:
    r = main.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.bold = False
    if sub:
        r.font.subscript = True
print("4.9.3 rebuilt")

# ---- 3. blank page: the section that starts with "1.2 Overview" began on a new page after a block of empty paragraphs
ps = d.paragraphs
ov = [i for i, p in enumerate(ps) if p.text.strip() == "Overview"][0]
for q in ps[ov - 6:ov]:
    if not q.text.strip() and not q._p.xpath(".//w:sectPr"):
        q._p.getparent().remove(q._p)
ps = d.paragraphs
ov = [i for i, p in enumerate(ps) if p.text.strip() == "Overview"][0]
# the sectPr that closes the Introduction's second section decides how that section starts
for q in ps[ov:ov + 40]:
    sp = q._p.xpath(".//w:sectPr")
    if sp:
        t = sp[0].find(qn("w:type"))
        if t is None:
            t = OxmlElement("w:type")
            sp[0].insert(0, t)
        t.set(qn("w:val"), "continuous")
        print("section after Overview made continuous")
        break
# and the section break paragraph just before Overview starts the section on the same page too
for q in ps[ov - 6:ov]:
    sp = q._p.xpath(".//w:sectPr")
    if sp:
        t = sp[0].find(qn("w:type"))
        if t is None:
            t = OxmlElement("w:type")
            sp[0].insert(0, t)
        t.set(qn("w:val"), "continuous")

# ---- 4. drop the appendices
ps = d.paragraphs
a = [i for i, p in enumerate(ps) if p.text.strip() == "APPENDICES"][0]
body = d.element.body
first = ps[a]._p
kill = []
node = first
while node is not None and node.tag != qn("w:sectPr"):
    kill.append(node)
    node = node.getnext()
for n in kill:
    n.getparent().remove(n)
print("appendices removed:", len(kill), "elements")
d.save(DOC)
print("saved", DOC)

# ---- 5. follow-up: line spacing of the rebuilt paragraph, and no empty trailing page
d = docx.Document(DOC)
ps = d.paragraphs
la = [p for p in ps if p.text.startswith("In the full-utility phase, each candidate")][0]
la.style = d.styles["Body Text"]
pf = la.paragraph_format
pf.left_indent, pf.first_line_indent, pf.line_spacing = 75565, 359410, 1.3
pf.space_before = pf.space_after = None
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
la.alignment = AL.JUSTIFY
n = 0
while ps and not ps[-1].text.strip() and not ps[-1]._p.xpath(".//w:sectPr") and not ps[-1]._p.xpath(".//w:drawing"):
    ps[-1]._p.getparent().remove(ps[-1]._p)
    ps = d.paragraphs
    n += 1
print("trailing empty paragraphs removed:", n)
d.save(DOC)

# ---- 6. merge the (now empty) last section into the bibliography section so no blank page is produced
d = docx.Document(DOC)
body = d.element.body
sect_ps = [p for p in d.paragraphs if p._p.xpath("./w:pPr/w:sectPr")]
last = sect_ps[-1]
sp = last._p.find(qn("w:pPr")).find(qn("w:sectPr"))
if not last.text.strip():
    body_sp = body.find(qn("w:sectPr"))
    last._p.find(qn("w:pPr")).remove(sp)
    body.replace(body_sp, sp)
    last._p.getparent().remove(last._p)
    print("last section merged")
d.save(DOC)

# ---- 7. Section 5.4 bullets that were split mid-sentence, and keep-with-next chains that left a nearly empty page
d = docx.Document(DOC)
ps = d.paragraphs


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


iso = [i for i, p in enumerate(ps) if p.text.startswith("Isolation. The memory keys")][0]
set_text(ps[iso], "Isolation. Memory keys are namespaced as scenario|planner|seed. An early sweep let the planners’ detection histories leak into each other "
                  "because identity matching was not scoped; a targeted test caught this, and it was fixed before any reported number was produced.")
for q in ps[iso + 1:iso + 4]:
    if q.text.strip() in ("planners’ detection histories are not crossed-checked so they overlap.", "produced.") or q.text.startswith("This is a targeted test"):
        q._p.getparent().remove(q._p)
ps = d.paragraphs
par = [p for p in ps if p.text.startswith("Parallel sweeps.")][0]
set_text(par, "Parallel sweeps. For the 31-wall yard each (seed, scenario) pair is a separate subprocess, and fourteen run at once "
              "(experiments/run_parallel_sweep.py); the parts are concatenated into a single CSV. Each large sweep is 240 rows and took about 30 minutes.")
cleared = 0
for p in d.paragraphs:
    if p.style.name in ("List Paragraph", "Body Text") and p.paragraph_format.keep_with_next and p.text.strip() and not p.text.startswith("Table "):
        p.paragraph_format.keep_with_next = None
        cleared += 1
print("keep-with-next cleared on", cleared, "paragraphs")
d.save(DOC)

# ---- 8. lookahead paragraph width, and the section break that left page 5.5 nearly empty
d = docx.Document(DOC)
ps = d.paragraphs
la = [p for p in ps if p.text.startswith("In the full-utility phase, each candidate")][0]
la.paragraph_format.right_indent = 685800
idx = [i for i, p in enumerate(ps) if p.text.startswith("Four defects in the work itself")][0]
brk = [i for i in range(idx, idx + 4) if ps[i]._p.xpath(".//w:sectPr")]
if brk:
    for j in range(brk[0] + 1, len(ps)):
        sp = ps[j]._p.xpath(".//w:sectPr")
        if sp:
            t = sp[0].find(qn("w:type"))
            if t is None:
                t = OxmlElement("w:type")
                sp[0].insert(0, t)
            t.set(qn("w:val"), "continuous")
            print("section after 5.5 intro made continuous")
            break
d.save(DOC)
