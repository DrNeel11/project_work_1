# -*- coding: utf-8 -*-
"""Formatting and consistency pass over the edited report (.docx). Input is the
already-aligned copy; output is UW-TIG_Project_Work_Report_final.docx.
  * body-text-styled-as-heading paragraphs restored to body text
  * no bold in body paragraphs; one font (Times New Roman) and size per kind of text
  * Table 2.1: empty column removed; text tables left-aligned; years dropped from planner names
  * Table 6.3: redundant no-uncertainty / no-growth rows removed, Section 6.4 rewritten to match
  * Section 6.1: unfair CPU-vs-GPU row removed from the detector comparison
  * Section 7 rewritten as one consistent list
  * Fig. 4.1 replaced by a black-and-white diagram
"""
import copy
import io
import os
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(ROOT, "UW-TIG_Project_Work_Report_edited_aligned.docx")
OUT = os.path.join(ROOT, "UW-TIG_Project_Work_Report_final.docx")
FIG = os.path.join(HERE, "fig41_architecture.png")
FONT = "Times New Roman"
BODY_PT, SMALL_PT = 12, 10.5

d = docx.Document(SRC)
P = lambda: d.paragraphs


def find(prefix, start=0):
    for i, p in enumerate(P()):
        if i >= start and p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def delete(p):
    p._p.getparent().remove(p._p)


def as_body(p):
    p.style = d.styles["Body Text"]
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent, pf.line_spacing = 75565, 359410, 1.3
    p.alignment = AL.JUSTIFY


# ------------------------------------------------------------------ 1. body text that was styled as a heading
for prefix in ("Rückin et al. [1] use a model", "\tUW-TIG reaches precision 0.949", "\tThe coverage guarantee is the key",
               "The research done by Rückin et al.", "We also tested if it"):
    for p in P():
        if p.text.startswith(prefix) and p.style.name.startswith("Heading"):
            as_body(p)
            if p.text.startswith("\t"):
                p.runs[0].text = p.runs[0].text.lstrip("\t")

# ------------------------------------------------------------------ 2. Section 6.1 (fair detector comparison)
p61 = find("The detector is evaluated on the held-out MBDD2025")
set_text(p61, "The detector is evaluated on the held-out MBDD2025 test split of 1,448 images. Two models were trained with identical settings "
              "(640 px, up to 100 epochs, the same L4 GPU and the same data): the default YOLOv8s, with the rarest class, bulge, oversampled "
              "about four times, and a larger YOLO11m retrain, which was evaluated but not adopted. An earlier YOLOv8n run (320 px, 30 epochs, "
              "CPU) was only a smoke test of the pipeline; its resolution, epochs and hardware differ, so it is not compared here. "
              "Larger YOLOv8 variants (m and l) were not trained.")
for t in d.tables:
    if t.rows[0].cells[0].text.strip() == "Model":
        for row in list(t.rows)[1:]:
            if row.cells[0].text.startswith("YOLOv8n"):
                row._tr.getparent().remove(row._tr)

# ------------------------------------------------------------------ 3. Table 6.3 + Section 6.4
for t in d.tables:
    if t.rows[0].cells[0].text.strip() == "Planner" and any(r.cells[0].text.startswith("no uncertainty term") for r in t.rows):
        for row in list(t.rows):
            if row.cells[0].text.startswith(("no uncertainty term", "no temporal (growth) term")):
                row._tr.getparent().remove(row._tr)
p63 = find("UW-TIG reaches precision 0.949")
set_text(p63, "UW-TIG reaches precision 0.949, recall 0.967 and F1 0.953. By scenario (precision / recall / F1) it scores static 0.964 / 0.933 / 0.940, "
              "uncertain 0.893 / 0.967 / 0.925 and multi-defect 0.989 / 1.000 / 0.995; the per-seed F1 values are 0.943, 0.890, 1.000, 1.000 and 0.933. "
              "Not every average is this good: the deliberately faint uncertain scenario has a precision of 0.893, and one test seed scores an F1 of only 0.890.")
p63b = find("The coverage guarantee is the key")
set_text(p63b, "The coverage guarantee is the component that makes the difference: without it recall falls from 0.967 to 0.556, F1 from 0.953 to 0.680, and "
               "coverage to 0.44 (4 of 9 walls). Removing staleness or lookahead leaves F1 at 0.958 (precision 0.967, recall 0.956), with a localisation "
               "error of 0.57–0.59 m and flights 32–36% shorter (80.6 m and 85.2 m against 125.1 m). The no-uncertainty and no-growth variants are "
               "omitted from Table 6.3 because they reproduce the full planner (Section 6.4).")
p64 = find("Table 6.3 shows that the no-uncertainty")
set_text(p64, "The no-uncertainty and no-growth ablations give the same precision, recall and F1 as the full planner (0.949, 0.967, 0.953); their flight "
              "distances are 124.6 m and 125.1 m against 125.1 m. They are therefore left out of Table 6.3, since their rows would repeat the first row. "
              "We checked the code rather than assume, and found three reasons.")
i64 = [i for i, p in enumerate(P()) if p.text.startswith("These terms are not covered")][0]
texts = [
    "The coverage phase never evaluates these terms. While any wall is unvisited, the planner scores only IG − 0.4·Cost. With nine walls and a "
    "16-viewpoint budget, roughly half of every mission is spent in this phase, and the three missions of a seed each repeat it because it is defined per mission.",
    "Growth is near zero by construction in these scenarios. Growth is the change in a defect’s measured size, clipped to be non-negative. Static, "
    "uncertain and multi-defect defects never grow, so the growth term only sees detector noise around a constant. Removing it leaves the flight distance "
    "identical to the decimal (125.1 m).",
    "Uncertainty is real but small. The TTA spread is typically 0.1–0.3, so weighting it by 1.5 gives a small term beside the others. Removing it "
    "shifts the path slightly (124.6 m against 125.1 m) but not the number of looks each wall receives, which is all that confirmed precision and recall depend on.",
]
for k, txt in enumerate(texts):
    set_text(P()[i64 + k], txt)

# ------------------------------------------------------------------ 4. Section 7
items = [
    "Make the novel terms matter. Re-weight or normalise the uncertainty and growth terms and test them on scenarios with real growth, so that the ablations in Section 6.4 separate.",
    "Better uncertainty. Try post-hoc temperature scaling first; then prototype MC-Dropout by adding dropout to the detection head and retraining.",
    "Close the growing-defect gap. Mix synthetic defect imagery into training, as synth-dacl does for bridge segmentation [29], or obtain a longitudinal real dataset. SDNET2018 was planned, but its download is blocked by a bot check.",
    "Use the broader detector in the loop. Texture the simulator with dacl10k-style close-ups, re-tune the reporting rule for the fine-tuned detector, and re-run Table 6.9.",
    "Lift the weak slices. Raise the uncertain scenario’s precision above 0.90 through scenario-specific tuning of the confirmation rule, and narrow Rückin’s precision lead by re-weighting the utility terms.",
    "Fuller baselines. Re-implement the compared papers, and IA-TIGRIS, as full systems with sampling-based planning and online retraining.",
    "More seeds for the ablations, and a decision on whether staleness and lookahead stay on by default.",
    "Richer simulation. Pipe segments and truss rigs as inspectable geometry, sensor noise and wind, and a Gaussian-splatting or game-engine renderer if the work moves to training on synthetic imagery.",
    "Larger detectors. Train YOLOv8m and YOLOv8l with the same settings as YOLOv8s (640 px, 100 epochs, same GPU) for a fair size comparison. The YOLO11m retrain (mAP50 0.890, precision 0.869, recall 0.863) was a marginal, mixed gain, and adopting it would need a full re-run of the sweeps. 1280 px was not tried.",
]
a = [i for i, p in enumerate(P()) if p.text.startswith("Make the new words count")][0]
for k, txt in enumerate(items):
    set_text(P()[a + k], txt)
# the old item 9 was split over three paragraphs; drop the two leftovers
for p in list(P())[a + len(items): a + len(items) + 2]:
    if p.style.name == "List Paragraph" and p.text.strip():
        delete(p)

# ------------------------------------------------------------------ 5. years out of planner names in tables
YEAR = re.compile(r"(\bet al\.|\(GATSBI\)|NBV|Isler|Bircher|Alamdari|Rückin)(?:,)?\s+(?:\(?(?:19|20)\d\d(?:[–\-](?:19|20)?\d\d)?\)?)")
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for q in c.paragraphs:
                new = YEAR.sub(lambda m: m.group(1), q.text)
                new = re.sub(r"\s*\((?:classical, )?\d{4}\)", lambda m: " (classical)" if "classical" in m.group(0) else "", new)
                if new != q.text and q.runs:
                    set_text(q, new)

# ------------------------------------------------------------------ 6. Table 2.1: remove the empty column
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
for t in d.tables:
    hdr = [c.text.strip() for c in t.rows[0].cells]
    if hdr[:3] == ["Paper and authors", "Technique", "Advantages"] and len(hdr) == 5:
        grid = t._tbl.tblGrid.findall(W + "gridCol")
        grid[3].getparent().remove(grid[3])
        for row in t.rows:
            tcs = row._tr.findall(W + "tc")
            tcs[3].getparent().remove(tcs[3])
        tot = sum(int(g.get(W + "w")) for g in t._tbl.tblGrid.findall(W + "gridCol"))
        want = int(6.14 * 1440)
        share = [0.27, 0.25, 0.23, 0.25]
        for g, s in zip(t._tbl.tblGrid.findall(W + "gridCol"), share):
            g.set(W + "w", str(int(want * s)))
        for row in t.rows:
            for tc, s in zip(row._tr.findall(W + "tc"), share):
                pr = tc.find(W + "tcPr")
                w = pr.find(W + "tcW") if pr is not None else None
                if w is not None:
                    w.set(W + "type", "dxa"); w.set(W + "w", str(int(want * s)))
        break

# ------------------------------------------------------------------ 7. body range, bold, fonts
start = end = None
for i, p in enumerate(P()):
    tx = p.text.strip()
    if start is None and i > 60 and tx == "INTRODUCTION" and p.style.name == "Heading 1":
        start = i
    if start and end is None and tx == "BIBLIOGRAPHY" and p.style.name == "Heading 1":
        end = i


def is_code(p):
    names = {r.font.name for r in p.runs if r.text.strip()}
    return "Courier New" in names


def set_font(run, size):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(att), FONT)
    if size:
        run.font.size = Pt(size)


counts = {"bold": 0, "font": 0}
for i, p in enumerate(P()):
    in_body = start <= i < end
    if p.style.name.startswith("Heading") or is_code(p):
        continue
    tx = p.text.strip()
    caption = tx.startswith(("Fig.", "Table "))
    for r in p.runs:
        if not r.text.strip():
            continue
        if r.font.name == "Geist Fallback":
            r.font.name = FONT
        if in_body:
            if r.bold and not caption:
                r.bold = False; counts["bold"] += 1
            set_font(r, SMALL_PT if caption else BODY_PT); counts["font"] += 1
        else:
            set_font(r, None)

for t in d.tables:
    text_table = sum(1 for c in t.rows[1].cells if re.search(r"[A-Za-z]{4}", c.text)) >= len(t.rows[1].cells) - 1
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            for q in c.paragraphs:
                if ri == 0:
                    q.alignment = AL.CENTER
                elif text_table:
                    q.alignment = AL.LEFT
                for r in q.runs:
                    set_font(r, SMALL_PT)
                    if ri > 0:
                        r.bold = False
print("bold runs cleared:", counts["bold"], "| runs normalised:", counts["font"])

# ------------------------------------------------------------------ 8. Fig. 4.1
cap = find("Fig. 4.1 System architecture")
idx = [i for i, p in enumerate(P()) if p.text.strip().startswith("Fig. 4.1 System architecture: the five-stage")][-1]
for j in range(idx - 4, idx):
    p = P()[j]
    blips = p._p.xpath(".//a:blip")
    if blips:
        rid = blips[0].get(qn("r:embed"))
        part = d.part.related_parts[rid]
        ext = os.path.splitext(part.partname)[1].lower()
        im = Image.open(FIG).convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG" if ext in (".jpg", ".jpeg") else "PNG", quality=95)
        part._blob = buf.getvalue()
        print("Fig 4.1 image replaced in", part.partname)
        break

# ------------------------------------------------------------------ 9. headings / keep-together (same as before)
for p in P():
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True
# captions stay with their tables: walk back over blank paragraphs from each table
for tbl in d.element.body.iter(qn("w:tbl")):
    prev = tbl.getprevious(); hops = 0
    while prev is not None and prev.tag == qn("w:p") and hops < 4:
        ppr = prev.find(qn("w:pPr"))
        if ppr is None:
            ppr = prev.makeelement(qn("w:pPr"), {}); prev.insert(0, ppr)
        if ppr.find(qn("w:keepNext")) is None:
            ppr.append(ppr.makeelement(qn("w:keepNext"), {}))
        txt = "".join(t.text or "" for t in prev.iter(qn("w:t"))).strip()
        if txt.startswith("Table "):
            break
        prev = prev.getprevious(); hops += 1

# a little more room between the running header and the first line of each page
from docx.shared import Inches
for sec in d.sections:
    sec.top_margin = sec.top_margin + Inches(0.12)
d.save(OUT)
print("saved", OUT)
